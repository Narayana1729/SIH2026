import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

for filename in ['SIH2026-IDEA-Presentation-Format.pptx', 'SIH2026_PyroSat_Submission.pptx']:
    prs = Presentation(filename)
    slide2 = prs.slides[1]

    DARK_BLUE = RGBColor(43, 91, 132)
    TEXT_BLACK = RGBColor(30, 30, 30)
    BOLD_BLACK = RGBColor(0, 0, 0)

    # Re-position or re-create the bottom two text boxes
    shapes_to_remove = []
    for s in slide2.shapes:
        if s.has_text_frame:
            txt = s.text_frame.text
            if "Proposed Solution" in txt or "Innovation" in txt:
                shapes_to_remove.append(s)

    for s in shapes_to_remove:
        sp = s._element
        sp.getparent().remove(sp)

    # 1. Left Box: Proposed Solution (Width ~ 4.25 inches)
    tb_sol = slide2.shapes.add_textbox(Inches(0.38), Inches(2.78), Inches(4.25), Inches(4.0))
    tf_sol = tb_sol.text_frame
    tf_sol.word_wrap = True
    tf_sol.margin_left = tf_sol.margin_right = tf_sol.margin_top = tf_sol.margin_bottom = 0

    p_hs = tf_sol.paragraphs[0]
    p_hs.text = "Proposed Solution :"
    p_hs.font.name = "Arial"
    p_hs.font.size = Pt(13)
    p_hs.font.bold = True
    p_hs.font.underline = True
    p_hs.font.color.rgb = DARK_BLUE
    p_hs.space_after = Pt(6)

    sol_bullets = [
        ("• When satellite telemetry arrives, the system applies ",
         "Dual-Band Planck Pyrometry",
         " to calculate true flame temperature (Tf > 1200 K) and exact fire area from sub-pixel clutter."),
        ("• A ",
         "2-Stage Hierarchical ML Classifier",
         " automatically segregates routine refinery flares from industrial explosions and forest wildfires."),
        ("• An integrated ",
         "3D GIS WebGL Command Dashboard",
         " drapes thermal hotspots and real-time Gaussian toxic plume evacuation zones onto a live virtual globe.")
    ]

    for prefix, bold_text, suffix in sol_bullets:
        p = tf_sol.add_paragraph()
        p.space_after = Pt(8)
        p.line_spacing = 1.15
        r = p.add_run(); r.text = prefix; r.font.name = "Arial"; r.font.size = Pt(9.5); r.font.color.rgb = TEXT_BLACK
        r = p.add_run(); r.text = bold_text; r.font.name = "Arial"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = BOLD_BLACK
        r = p.add_run(); r.text = suffix; r.font.name = "Arial"; r.font.size = Pt(9.5); r.font.color.rgb = TEXT_BLACK

    # 2. Right Box: Innovation / Uniqueness (Width ~ 4.25 inches, Left = 8.70 inches)
    # This leaves the center space from 4.70 in to 8.65 in (almost 4 inches wide!) completely open for the image!
    tb_inn = slide2.shapes.add_textbox(Inches(8.70), Inches(2.78), Inches(4.25), Inches(4.0))
    tf_inn = tb_inn.text_frame
    tf_inn.word_wrap = True
    tf_inn.margin_left = tf_inn.margin_right = tf_inn.margin_top = tf_inn.margin_bottom = 0

    p_hi = tf_inn.paragraphs[0]
    p_hi.text = "Innovation/Uniqueness:"
    p_hi.font.name = "Arial"
    p_hi.font.size = Pt(13)
    p_hi.font.bold = True
    p_hi.font.underline = True
    p_hi.font.color.rgb = DARK_BLUE
    p_hi.space_after = Pt(6)

    inn_bullets = [
        ("• ", "India's first sub-pixel pyrometry platform", " to detect industrial fires."),
        ("• ", "Converts raw heat pixels into physical flame data", " (temperature & burning area) with zero AI hallucination."),
        ("• ", "Dual-satellite land grounding", " uses 10m ESA land-cover and Sentinel-2 data to eliminate false alarms at industrial plants."),
        ("• ", "Live Toxic Plume Dispersion Engine", " combines real-time wind and chemical data to map evacuation zones for first responders.")
    ]

    for prefix, bold_text, suffix in inn_bullets:
        p = tf_inn.add_paragraph()
        p.space_after = Pt(7)
        p.line_spacing = 1.15
        r = p.add_run(); r.text = prefix; r.font.name = "Arial"; r.font.size = Pt(9.5); r.font.color.rgb = TEXT_BLACK
        r = p.add_run(); r.text = bold_text; r.font.name = "Arial"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = BOLD_BLACK
        r = p.add_run(); r.text = suffix; r.font.name = "Arial"; r.font.size = Pt(9.5); r.font.color.rgb = TEXT_BLACK

    prs.save(filename)
    print(f"Successfully added center space on Slide 2 for {filename}!")
