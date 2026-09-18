import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

for filename in ['SIH2026-IDEA-Presentation-Format.pptx', 'SIH2026_PyroSat_Submission.pptx']:
    prs = Presentation(filename)
    slide2 = prs.slides[1]

    DARK_BLUE = RGBColor(43, 91, 132)
    TEXT_BLACK = RGBColor(30, 30, 30)
    BOLD_BLACK = RGBColor(0, 0, 0)

    # Find the Innovation textbox
    for s in slide2.shapes:
        if s.has_text_frame and "Innovation" in s.text_frame.text:
            tf = s.text_frame
            tf.clear()

            p_hi = tf.paragraphs[0]
            p_hi.text = "Innovation/Uniqueness:"
            p_hi.font.name = "Arial"
            p_hi.font.size = Pt(13)
            p_hi.font.bold = True
            p_hi.font.underline = True
            p_hi.font.color.rgb = DARK_BLUE
            p_hi.space_after = Pt(6)

            inn_bullets = [
                ("• ", "India's first sub-pixel pyrometry system", " to detect industrial fires."),
                ("• ", "Converts raw heat pixels into physical flame data", " (temperature & burning area) with zero AI hallucination."),
                ("• ", "Dual-satellite land grounding", " uses 10m ESA land-cover and Sentinel-2 data to eliminate false alarms at industrial plants."),
                ("• ", "Live Toxic Plume Dispersion Engine", " combines real-time wind and chemical data to map evacuation zones for first responders.")
            ]

            for prefix, bold_text, suffix in inn_bullets:
                p = tf.add_paragraph()
                p.space_after = Pt(8)
                p.line_spacing = 1.15
                
                r = p.add_run(); r.text = prefix; r.font.name = "Arial"; r.font.size = Pt(10); r.font.color.rgb = TEXT_BLACK
                r = p.add_run(); r.text = bold_text; r.font.name = "Arial"; r.font.size = Pt(10); r.font.bold = True; r.font.color.rgb = BOLD_BLACK
                r = p.add_run(); r.text = suffix; r.font.name = "Arial"; r.font.size = Pt(10); r.font.color.rgb = TEXT_BLACK

    prs.save(filename)
    print(f"Updated {filename} successfully!")
