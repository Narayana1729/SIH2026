import os
from PIL import Image
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def find_asset(name):
    script_dir = os.path.dirname(os.path.abspath(__file__))
    candidates = [
        name,
        os.path.join('presentation', 'assets', name),
        os.path.join(script_dir, '..', 'assets', name),
        os.path.join(script_dir, name)
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return name

def find_ppt(name):
    script_dir = os.path.dirname(os.path.abspath(__file__))
    candidates = [
        name,
        os.path.join(script_dir, '..', '..', name),
        os.path.join(script_dir, '..', name),
        os.path.join(script_dir, name)
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return name

def prepare_clean_hex_background():
    tpl_path = find_asset('seven_hexagons_template.png')
    im = Image.open(tpl_path)
    # Hexagons span X:[40, 985], Y:[195, 355]
    crop_hex = im.crop((40, 195, 985, 355))
    out_dir = os.path.dirname(tpl_path)
    out_path = os.path.join(out_dir, 'clean_hex_bg.png')
    crop_hex.save(out_path)
    print(f"Prepared clean_hex_bg at {out_path}")
    return out_path

def build_real_ppt_slide6():
    hex_bg_img = prepare_clean_hex_background()

    HEADER_LINE = RGBColor(70, 130, 180)     # Steel blue
    LINE_GRAY = RGBColor(203, 213, 225)      # #CBD5E1
    NAVY_TITLE = RGBColor(15, 41, 66)        # #0F2942
    TEXT_MUTED = RGBColor(71, 85, 105)       # #475569
    TEXT_BLACK = RGBColor(40, 45, 55)
    BOLD_BLACK = RGBColor(0, 0, 0)
    LINK_BLUE = RGBColor(14, 75, 160)        # Professional navy/hyperlink blue

    # 7 References data matching user's exact order, concise copy, and official hyperlinks
    references_data = [
        {
            "title_l1": "NASA FIRMS",
            "title_l2": "Telemetry",
            "desc": "NRT Thermal Hotspot Ingestion — VIIRS 375m & MODIS 1km active fire feeds.",
            "url": "https://firms.modaps.eosdis.nasa.gov/"
        },
        {
            "title_l1": "Dozier Pyrometry",
            "title_l2": "(1981)",
            "desc": "Sub-Pixel Planck Inversion — Resolves flame temp (Tf > 1100 K) & sub-pixel area (Af).",
            "url": "https://doi.org/10.1016/0034-4257(81)90021-3"
        },
        {
            "title_l1": "ESA WorldCover",
            "title_l2": "10m",
            "desc": "10m Surface Ground Truth — Sentinel-1/2 LULC classifying Built-up (Class 50) vs Cropland.",
            "url": "https://esa-worldcover.org/"
        },
        {
            "title_l1": "GEM & World Bank",
            "title_l2": "GGFR",
            "desc": "2,000+ Facility Registry — Catalog of Indian refineries, petrochemical hubs & flaring stacks.",
            "url": "https://globalenergymonitor.org/"
        },
        {
            "title_l1": "NOAA CAMEO &",
            "title_l2": "NIOSH",
            "desc": "Chemical & HazMat Profiles — Reactivity data, UN placards & toxic thresholds.",
            "url": "https://cameochemicals.noaa.gov/"
        },
        {
            "title_l1": "EPA / Briggs",
            "title_l2": "Plume Rise",
            "desc": "Atmospheric Dispersion — Buoyant plume rise & Pasquill-Gifford Gaussian dispersion.",
            "url": "https://www.epa.gov/scram/air-quality-dispersion-modeling"
        },
        {
            "title_l1": "NDRF & HazMat",
            "title_l2": "ERG",
            "desc": "Tactical Standoff & Routing — Emergency chemical standoff, isolation corridors & automated dispatch.",
            "url": "https://www.phmsa.dot.gov/erg"
        }
    ]

    for ppt_base in ['SIH2026_PyroSat_Submission.pptx', 'SIH2026-IDEA-Presentation-Format.pptx']:
        ppt_path = find_ppt(ppt_base)
        if not os.path.exists(ppt_path):
            continue
        prs = Presentation(ppt_path)
        slide6 = prs.slides[5]

        # 1. Clean all shapes on slide 6 except template header items
        shapes_to_delete = []
        for s in slide6.shapes:
            if s.name in ["Rectangle 8", "Rectangle 9", "Slide Number Placeholder 5", "Footer Placeholder 6", "Oval 9", "Picture 10"]:
                continue
            shapes_to_delete.append(s)

        for s in shapes_to_delete:
            sp = s._element
            sp.getparent().remove(sp)

        # 2. Add Clean Slide Header Title
        tb_title = slide6.shapes.add_textbox(Inches(1.8), Inches(0.12), Inches(8.8), Inches(1.0))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = "RESEARCH AND REFERENCES"
        p_title.alignment = PP_ALIGN.CENTER
        p_title.font.name = "Times New Roman"
        p_title.font.size = Pt(18)
        p_title.font.bold = True
        p_title.font.color.rgb = BOLD_BLACK

        # Underline under header
        l_top = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.28), Inches(1.18), Inches(12.75), Inches(0.02))
        l_top.fill.solid(); l_top.fill.fore_color.rgb = HEADER_LINE; l_top.line.fill.background()

        # -------------------------------------------------------------
        # 3. TOP SECTION: 7 HEXAGONS (Image Background + Hyperlinked Text)
        # -------------------------------------------------------------
        hex_row_left = Inches(0.38)
        hex_row_top = Inches(1.48)
        hex_row_width = Inches(12.56)
        hex_row_height = Inches(2.10)

        # Add the 7 hexagons image as the clean background
        slide6.shapes.add_picture(hex_bg_img, hex_row_left, hex_row_top, width=hex_row_width, height=hex_row_height)

        # Place real, editable text boxes over each hexagon with real hyperlinks
        hex_step = 12.56 / 7.0  # ~1.794 inches per hexagon
        for i, ref in enumerate(references_data):
            h_left = Inches(0.38 + i * hex_step)

            # Editable Text Frame inside hexagon body (below icon)
            tb_body = slide6.shapes.add_textbox(h_left + Inches(0.10), Inches(2.18), Inches(1.58), Inches(1.36))
            tf_body = tb_body.text_frame
            tf_body.word_wrap = True
            tf_body.margin_left = tf_body.margin_right = tf_body.margin_top = tf_body.margin_bottom = 0

            # Title Paragraph with active clickable hyperlink
            p_head = tf_body.paragraphs[0]
            p_head.alignment = PP_ALIGN.CENTER
            p_head.space_after = Pt(3)

            # Line 1 of title
            r1 = p_head.add_run()
            r1.text = ref["title_l1"] + "\n"
            r1.font.name = "Arial"
            r1.font.size = Pt(9.5)
            r1.font.bold = True
            r1.font.underline = True
            r1.font.color.rgb = LINK_BLUE
            r1.hyperlink.address = ref["url"]

            # Line 2 of title
            r2 = p_head.add_run()
            r2.text = ref["title_l2"]
            r2.font.name = "Arial"
            r2.font.size = Pt(9.5)
            r2.font.bold = True
            r2.font.underline = True
            r2.font.color.rgb = LINK_BLUE
            r2.hyperlink.address = ref["url"]

            # Description Paragraph (Regular, non-squished, clean typography)
            p_desc = tf_body.add_paragraph()
            p_desc.text = ref["desc"]
            p_desc.alignment = PP_ALIGN.CENTER
            p_desc.font.name = "Arial"
            p_desc.font.size = Pt(7.5)
            p_desc.font.color.rgb = TEXT_BLACK
            p_desc.space_before = Pt(2)

        # -------------------------------------------------------------
        # 4. MIDDLE DIVIDER & "UI/UX" BADGE
        # -------------------------------------------------------------
        mid_y = Inches(3.72)
        l_mid = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.38), mid_y, Inches(11.20), Inches(0.015))
        l_mid.fill.solid(); l_mid.fill.fore_color.rgb = LINE_GRAY; l_mid.line.fill.background()

        # UI/UX text badge on right
        tb_ui = slide6.shapes.add_textbox(Inches(11.75), Inches(3.52), Inches(1.20), Inches(0.35))
        tf_ui = tb_ui.text_frame; tf_ui.margin_left = tf_ui.margin_right = tf_ui.margin_top = tf_ui.margin_bottom = 0
        p_ui = tf_ui.paragraphs[0]
        p_ui.text = "UI/UX"
        p_ui.font.name = "Arial"
        p_ui.font.size = Pt(15)
        p_ui.font.bold = True
        p_ui.font.color.rgb = RGBColor(148, 163, 184)

        # -------------------------------------------------------------
        # 5. BOTTOM SECTION: 3 REAL UI/UX PICTURES & EDITABLE TEXT CAPTIONS
        # -------------------------------------------------------------
        bottom_top = Inches(3.88)
        pic_height = Inches(2.35)
        card_width = Inches(3.96)
        card_gap = Inches(0.34)

        ui_cards = [
            {
                "img": find_asset("cesium_pyrosat_hud_globe.png"),
                "left": Inches(0.38),
                "title": "3D Geospatial HUD (Cesium WebGL)",
                "title_color": RGBColor(2, 132, 199),
                "desc": "Real-time thermal clustering across Indian industrial corridors."
            },
            {
                "img": find_asset("gaussian_plume_hazard_rings.jpg"),
                "left": Inches(0.38 + card_width + card_gap),
                "title": "3D Toxic Plume & Evacuation Rings",
                "title_color": RGBColor(16, 185, 129),
                "desc": "Atmospheric dispersion driven by live Open-Meteo wind vectors."
            },
            {
                "img": find_asset("sanitized_dispatch_modal.png"),
                "left": Inches(0.38 + 2 * (card_width + card_gap)),
                "title": "CAMEO HazMat Dossier & NDRF Dispatch",
                "title_color": RGBColor(217, 119, 6),
                "desc": "Automated Incident Action Plan with SMS/WhatsApp responder routing."
            }
        ]

        for card in ui_cards:
            if os.path.exists(card["img"]):
                slide6.shapes.add_picture(card["img"], card["left"], bottom_top, width=card_width, height=pic_height)

            # Editable Caption Textbox Underneath Picture
            tb_cap = slide6.shapes.add_textbox(card["left"], bottom_top + pic_height + Inches(0.06), card_width, Inches(0.65))
            tf_cap = tb_cap.text_frame
            tf_cap.word_wrap = True
            tf_cap.margin_left = tf_cap.margin_right = tf_cap.margin_top = tf_cap.margin_bottom = 0

            p_t = tf_cap.paragraphs[0]
            p_t.text = card["title"]
            p_t.font.name = "Arial"
            p_t.font.size = Pt(9.5)
            p_t.font.bold = True
            p_t.font.color.rgb = card["title_color"]
            p_t.space_after = Pt(2)

            p_d = tf_cap.add_paragraph()
            p_d.text = card["desc"]
            p_d.font.name = "Arial"
            p_d.font.size = Pt(8)
            p_d.font.color.rgb = TEXT_MUTED

        prs.save(ppt_path)
        print(f"Updated {ppt_path} Slide 6 with live hyperlinks and balanced typography!")

if __name__ == '__main__':
    build_real_ppt_slide6()
