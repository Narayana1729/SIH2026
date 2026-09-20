import os
from PIL import Image, ImageDraw, ImageFont
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_filled_hexagons_image(output_path="seven_references_filled.png"):
    # Load user's cropped 7 hexagons
    base = Image.open('scaled_hexagons.png').convert('RGB')
    draw = ImageDraw.Draw(base)

    f_bold = '/System/Library/Fonts/Supplemental/Arial Bold.ttf'
    f_reg = '/System/Library/Fonts/Supplemental/Arial.ttf'
    if not os.path.exists(f_bold):
        f_bold = '/System/Library/Fonts/Helvetica.ttc'
        f_reg = '/System/Library/Fonts/Helvetica.ttc'

    font_title = ImageFont.truetype(f_bold, 21)
    font_body = ImageFont.truetype(f_reg, 13.5)
    font_link = ImageFont.truetype(f_bold, 13)

    refs_data = [
        {
            'title': 'NASA FIRMS',
            'sub': 'VIIRS (375m) & MODIS\nNRT fire telemetry\n(I-Band radiance)',
            'link': 'LINK'
        },
        {
            'title': 'GEM Registry',
            'sub': '2,000+ Indian plants\nRefineries, steel &\nGGFR flare stacks',
            'link': 'LINK'
        },
        {
            'title': 'Dozier Planck',
            'sub': 'Dual-band pyrometry\nResolves true flame\nTf (>1100K) & Af',
            'link': 'LINK'
        },
        {
            'title': 'NOAA CAMEO',
            'sub': 'Chemical toxicity\nUN HazMat profiles\nERPG / IDLH limits',
            'link': 'LINK'
        },
        {
            'title': 'ESA 10m LULC',
            'sub': 'Sentinel-1/2 LULC\nVerifies built-up\nClass 50 vs crops',
            'link': 'LINK'
        },
        {
            'title': 'Briggs Plume',
            'sub': 'Atmospheric physics\nGaussian dispersion\nLive wind vectors',
            'link': 'LINK'
        },
        {
            'title': 'NDRF & ERG',
            'sub': 'Tactical standoff\nIsolation corridors\nAutomated dispatch',
            'link': 'LINK'
        }
    ]

    W, H = base.size
    x_step = W / 7.0

    for i, ref in enumerate(refs_data):
        cx = int((i + 0.5) * x_step)
        
        # Title
        draw.text((cx, 138), ref['title'], font=font_title, fill='#0F2942', anchor='mt')
        draw.line([(cx - 70, 168), (cx + 70, 168)], fill='#CBD5E1', width=1)
        
        # Description
        lines = ref['sub'].split('\n')
        for l_idx, line in enumerate(lines):
            draw.text((cx, 178 + l_idx * 21), line, font=font_body, fill='#334155', anchor='mt')
            
        # Link pill
        pw, ph = 76, 26
        px = cx - pw // 2
        py = 312
        draw.rounded_rectangle([(px, py), (px + pw, py + ph)], radius=8, fill='#1D4ED8')
        draw.text((cx, py + ph // 2), ref['link'], font=font_link, fill='#FFFFFF', anchor='mm')

    base.save(output_path)
    print(f"Filled hexagons saved to {output_path}")
    return output_path

def generate_full_slide6_image(hex_img_path, output_path="pyrosat_slide6_references_filled.png"):
    W, H = 2560, 1440
    canvas = Image.new('RGB', (W, H), color='#FFFFFF')
    draw = ImageDraw.Draw(canvas)

    f_bold = '/System/Library/Fonts/Supplemental/Arial Bold.ttf'
    f_reg = '/System/Library/Fonts/Supplemental/Arial.ttf'
    f_serif = '/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf'
    if not os.path.exists(f_bold):
        f_bold = '/System/Library/Fonts/Helvetica.ttc'
        f_reg = '/System/Library/Fonts/Helvetica.ttc'
        f_serif = f_bold

    font_main_title = ImageFont.truetype(f_serif, 52)
    font_team = ImageFont.truetype(f_bold, 24)
    font_sih = ImageFont.truetype(f_bold, 26)
    font_top_note = ImageFont.truetype(f_bold, 14)
    font_uiux = ImageFont.truetype(f_bold, 42)
    font_ui_title = ImageFont.truetype(f_bold, 18)
    font_ui_sub = ImageFont.truetype(f_reg, 13)
    font_footer = ImageFont.truetype(f_reg, 16)

    # 1. Top Header
    draw.rounded_rectangle([(70, 38), (340, 114)], radius=38, outline='#0284C7', width=3, fill='#F0F9FF')
    draw.text((120, 50), "SIH 2026", font=font_team, fill='#0369A1')
    draw.text((110, 78), "Team Thinkers", font=ImageFont.truetype(f_reg, 17), fill='#0284C7')

    draw.text((W // 2, 50), "RESEARCH AND REFERENCES", font=font_main_title, fill='#1E3A8A', anchor="mt")

    draw.rounded_rectangle([(W - 380, 38), (W - 70, 118)], radius=12, outline='#E2E8F0', width=1, fill='#F8FAFC')
    draw.text((W - 360, 46), "SMART INDIA", font=font_sih, fill='#EA580C')
    draw.text((W - 360, 76), "HACKATHON 2026", font=font_sih, fill='#15803D')

    draw.line([(70, 132), (W - 70, 132)], fill='#CBD5E1', width=2)

    # 2. Top Notes above each hexagon
    top_notes = [
        "Telemetry Ingestion",
        "Facility Registry",
        "Planck Pyrometry",
        "Chemical Database",
        "Surface Ground Truth",
        "Plume Dispersion",
        "Tactical Dispatch"
    ]
    hex_w = 2362
    x_offset = (W - hex_w) // 2
    x_step = hex_w / 7.0

    for i, note in enumerate(top_notes):
        cx = int(x_offset + (i + 0.5) * x_step)
        draw.text((cx, 146), note, font=font_top_note, fill='#64748B', anchor="mt")

    # 3. Paste the 7 Filled Hexagons
    hex_im = Image.open(hex_img_path)
    canvas.paste(hex_im, (x_offset, 178))

    # 4. Middle Divider & "UI/UX" Badge
    mid_y = 605
    draw.line([(70, mid_y), (W - 250, mid_y)], fill='#CBD5E1', width=2)
    draw.text((W - 170, mid_y - 25), "UI/UX", font=font_uiux, fill='#94A3B8')

    # Curved pointer arrow
    draw.arc([(W - 220, mid_y + 15), (W - 140, mid_y + 95)], 90, 270, fill='#94A3B8', width=3)
    draw.polygon([(W - 180, mid_y + 98), (W - 192, mid_y + 88), (W - 192, mid_y + 108)], fill='#94A3B8')

    # 5. Bottom Half: 3 UI/UX Prototype Screenshots
    img_globe_path = 'cesium_pyrosat_hud_globe.png'
    img_plume_path = 'gaussian_plume_hazard_rings.jpg'
    img_dispatch_path = 'sanitized_dispatch_modal.png'

    screen_y = 645
    screen_h = 715

    # Screen 1 (Left): 3D Globe HUD
    c1_x, c1_w = 70, 750
    draw.rounded_rectangle([(c1_x, screen_y), (c1_x + c1_w, screen_y + screen_h)], radius=16, fill='#0F172A', outline='#CBD5E1', width=2)
    if os.path.exists(img_globe_path):
        im_g = Image.open(img_globe_path).convert('RGB')
        im_g = im_g.resize((c1_w - 12, screen_h - 85), Image.Resampling.LANCZOS)
        canvas.paste(im_g, (c1_x + 6, screen_y + 6))
    draw.text((c1_x + 25, screen_y + screen_h - 60), "LIVE 3D GEOSPATIAL HUD (CESIUM WEBGL)", font=font_ui_title, fill='#38BDF8')
    draw.text((c1_x + 25, screen_y + screen_h - 35), "Full-globe thermal anomaly clustering across major Indian industrial corridors.", font=font_ui_sub, fill='#94A3B8')

    # Screen 2 (Center): 3D Gaussian Plume
    c2_x, c2_w = 860, 780
    draw.rounded_rectangle([(c2_x, screen_y), (c2_x + c2_w, screen_y + screen_h)], radius=16, fill='#0F172A', outline='#10B981', width=2)
    if os.path.exists(img_plume_path):
        im_p = Image.open(img_plume_path).convert('RGB')
        im_p = im_p.resize((c2_w - 12, screen_h - 85), Image.Resampling.LANCZOS)
        canvas.paste(im_p, (c2_x + 6, screen_y + 6))
    draw.text((c2_x + 25, screen_y + screen_h - 60), "3D TOXIC PLUME DISPERSION & EVACUATION RINGS", font=font_ui_title, fill='#34D399')
    draw.text((c2_x + 25, screen_y + screen_h - 35), "Dynamic atmospheric dispersion coupled with live Open-Meteo wind vectors.", font=font_ui_sub, fill='#94A3B8')

    # Screen 3 (Right): CAMEO HazMat Dossier & Dispatch
    c3_x, c3_w = 1680, 810
    draw.rounded_rectangle([(c3_x, screen_y), (c3_x + c3_w, screen_y + screen_h)], radius=16, fill='#0F172A', outline='#F59E0B', width=2)
    if os.path.exists(img_dispatch_path):
        im_d = Image.open(img_dispatch_path).convert('RGB')
        im_d = im_d.resize((c3_w - 12, screen_h - 85), Image.Resampling.LANCZOS)
        canvas.paste(im_d, (c3_x + 6, screen_y + 6))
    draw.text((c3_x + 25, screen_y + screen_h - 60), "CAMEO CHEMICAL DOSSIER & NDRF DISPATCH", font=font_ui_title, fill='#FBBF24')
    draw.text((c3_x + 25, screen_y + screen_h - 35), "Automated multi-agency Incident Action Plan (IAP) with chemical standoff directives.", font=font_ui_sub, fill='#94A3B8')

    # Footer
    draw.rectangle([(0, H - 45), (W, H)], fill='#0284C7')
    draw.text((70, H - 34), "@SIH Idea submission — PyroSat Satellite Thermal Intelligence System", 
              font=font_footer, fill='#FFFFFF')
    draw.text((W - 120, H - 34), "Slide 6", font=ImageFont.truetype(f_bold, 18), fill='#FFFFFF')

    canvas.save(output_path, quality=95)
    print(f"Full Slide 6 Image successfully generated: {output_path}")
    return output_path

def update_pptx_slide6(full_slide_img):
    for ppt_name in ['SIH2026_PyroSat_Submission.pptx', 'SIH2026-IDEA-Presentation-Format.pptx']:
        if not os.path.exists(ppt_name):
            continue
        prs = Presentation(ppt_name)
        slide6 = prs.slides[5]

        # Clean dynamic shapes on slide 6
        shapes_to_delete = []
        for s in slide6.shapes:
            if s.name in ["Rectangle 8", "Rectangle 9", "Slide Number Placeholder 5", "Footer Placeholder 6", "Oval 9", "Picture 10"]:
                continue
            shapes_to_delete.append(s)

        for s in shapes_to_delete:
            sp = s._element
            sp.getparent().remove(sp)

        # Add Full Slide Background Image matching the exact reference layout
        slide6.shapes.add_picture(full_slide_img, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))

        prs.save(ppt_name)
        print(f"Updated {ppt_name} Slide 6 successfully!")

if __name__ == '__main__':
    hex_img = build_filled_hexagons_image()
    full_slide = generate_full_slide6_image(hex_img)
    update_pptx_slide6(full_slide)
