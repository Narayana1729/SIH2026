import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def find_asset(name):
    script_dir = os.path.dirname(os.path.abspath(__file__))
    candidates = [
        name,
        os.path.join('presentation', 'assets', name),
        os.path.join(script_dir, '..', 'assets', name),
        os.path.join(script_dir, name)
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return name

def find_ppt(name):
    script_dir = os.path.dirname(os.path.abspath(__file__))
    candidates = [
        name,
        os.path.join(script_dir, '..', '..', name),
        os.path.join(script_dir, '..', name),
        os.path.join(script_dir, name)
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return name

def build_slide2_idea_title():
    # Color Palette: Professional ATS-Clean Enterprise Palette
    DARK_NAVY = RGBColor(15, 23, 42)        # #0F172A
    CARD_DARK = RGBColor(30, 41, 59)        # #1E293B
    BORDER_DARK = RGBColor(51, 65, 85)      # #334155
    AMBER_GOLD = RGBColor(245, 158, 11)     # #F59E0B
    CYAN_BLUE = RGBColor(56, 189, 248)      # #38BDF8
    WHITE = RGBColor(248, 250, 252)         # #F8FAFC
    HEADER_LINE = RGBColor(70, 130, 180)    # Steel blue
    LINE_GRAY = RGBColor(226, 232, 240)     # #E2E8F0
    TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B
    TEXT_CHARCOAL = RGBColor(51, 65, 85)    # #334155
    LINK_BLUE = RGBColor(2, 132, 199)       # #0284C7

    # Risk vs Solution Colors
    RED_BG = RGBColor(254, 226, 226)        # #FEE2E2
    RED_BORDER = RGBColor(239, 68, 68)      # #EF4444
    RED_TEXT = RGBColor(185, 28, 28)        # #B91C1C

    GREEN_BG = RGBColor(220, 252, 231)      # #DCFCE7
    GREEN_BORDER = RGBColor(34, 197, 94)    # #22C55E
    GREEN_TEXT = RGBColor(21, 128, 61)      # #15803D

    for ppt_base in ['SIH2026_PyroSat_Submission.pptx', 'SIH2026-IDEA-Presentation-Format.pptx']:
        ppt_path = find_ppt(ppt_base)
        if not os.path.exists(ppt_path):
            continue
        prs = Presentation(ppt_path)
        slide2 = prs.slides[1]

        # 1. Clean previous content shapes (preserve template header elements)
        shapes_to_delete = []
        for s in slide2.shapes:
            if s.name in ["Rectangle 8", "Slide Number Placeholder 5", "Footer Placeholder 6", "Oval 9", "Picture 10"]:
                continue
            shapes_to_delete.append(s)

        for s in shapes_to_delete:
            sp = s._element
            sp.getparent().remove(sp)

        # 2. Header Section: Idea Title & Tagline
        tb_title = slide2.shapes.add_textbox(Inches(1.8), Inches(0.10), Inches(8.8), Inches(1.02))
        tf_title = tb_title.text_frame
        tf_title.word_wrap = True
        tf_title.margin_top = tf_title.margin_bottom = tf_title.margin_left = tf_title.margin_right = 0
        
        p_title1 = tf_title.paragraphs[0]
        p_title1.text = "PyroSat: AI-Powered Satellite Thermal Intelligence"
        p_title1.alignment = PP_ALIGN.CENTER
        p_title1.font.name = "Arial"
        p_title1.font.size = Pt(17)
        p_title1.font.bold = True
        p_title1.font.color.rgb = DARK_NAVY
        p_title1.space_after = Pt(2)

        p_title2 = tf_title.add_paragraph()
        p_title2.text = "Precision Sub-Pixel Pyrometry, Industrial Fire Segregation & Real-Time Toxic Plume Defense for India"
        p_title2.alignment = PP_ALIGN.CENTER
        p_title2.font.name = "Arial"
        p_title2.font.size = Pt(9.5)
        p_title2.font.bold = True
        p_title2.font.color.rgb = RGBColor(71, 85, 105)

        # Underline divider under header
        l_top = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.38), Inches(1.18), Inches(12.56), Inches(0.02))
        l_top.fill.solid(); l_top.fill.fore_color.rgb = HEADER_LINE; l_top.line.fill.background()

        # ---------------------------------------------------------------------
        # 3. LEFT COLUMN: 3 High-Impact Cards + Working Prototype Badge
        # ---------------------------------------------------------------------
        left_x = Inches(0.38)
        left_w = Inches(3.80)
        card_start_y = Inches(1.28)

        left_cards_data = [
            {
                "tag": "Real-World Issue:",
                "tag_color": AMBER_GOLD,
                "text": "Coarse 375m satellite thermal pixels cannot distinguish routine industrial flaring from runaway refinery explosions or crop stubble, causing severe operational blindness.",
                "h": Inches(1.18)
            },
            {
                "tag": "Why Important:",
                "tag_color": AMBER_GOLD,
                "text": "India houses 2,000+ hazardous petrochemical and industrial facilities near dense populations. Delayed hazard response risks thousands of lives and ₹10,000+ Cr in critical assets.",
                "h": Inches(1.18)
            },
            {
                "tag": "Proposed Solution:",
                "tag_color": CYAN_BLUE,
                "text": "PyroSat solves sub-pixel flame temperature (Tf > 1100 K) and burning area via Planck pyrometry, eliminates false alarms using 10m ESA ground truth, and projects live 3D evacuation zones.",
                "h": Inches(1.25)
            }
        ]

        curr_y = card_start_y
        for c in left_cards_data:
            # Rounded card background
            card_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, curr_y, left_w, c["h"])
            card_bg.fill.solid(); card_bg.fill.fore_color.rgb = CARD_DARK
            card_bg.line.color.rgb = BORDER_DARK
            card_bg.line.width = Pt(1)

            tb_c = slide2.shapes.add_textbox(left_x + Inches(0.12), curr_y + Inches(0.08), left_w - Inches(0.24), c["h"] - Inches(0.14))
            tf_c = tb_c.text_frame
            tf_c.word_wrap = True
            tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0

            p_t = tf_c.paragraphs[0]
            p_t.text = c["tag"]
            p_t.font.name = "Arial"
            p_t.font.size = Pt(9.5)
            p_t.font.bold = True
            p_t.font.color.rgb = c["tag_color"]
            p_t.space_after = Pt(2)

            p_b = tf_c.add_paragraph()
            p_b.text = c["text"]
            p_b.font.name = "Arial"
            p_b.font.size = Pt(7.8)
            p_b.font.color.rgb = WHITE

            curr_y += c["h"] + Inches(0.12)

        # Working Prototype Callout Card
        proto_h = Inches(1.50)
        proto_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, curr_y, left_w, proto_h)
        proto_bg.fill.solid(); proto_bg.fill.fore_color.rgb = RGBColor(240, 249, 255)
        proto_bg.line.color.rgb = RGBColor(186, 230, 253)
        proto_bg.line.width = Pt(1.2)

        tb_proto = slide2.shapes.add_textbox(left_x + Inches(0.14), curr_y + Inches(0.08), left_w - Inches(0.28), proto_h - Inches(0.14))
        tf_proto = tb_proto.text_frame
        tf_proto.word_wrap = True
        tf_proto.margin_left = tf_proto.margin_right = tf_proto.margin_top = tf_proto.margin_bottom = 0

        p_ph = tf_proto.paragraphs[0]
        p_ph.text = "⚡ Working Prototype & Demonstrations:"
        p_ph.font.name = "Arial"
        p_ph.font.size = Pt(9)
        p_ph.font.bold = True
        p_ph.font.color.rgb = RGBColor(3, 105, 161)
        p_ph.space_after = Pt(3)

        p_pl1 = tf_proto.add_paragraph()
        r_pl1 = p_pl1.add_run()
        r_pl1.text = "🔗 Live Cloud Platform: pyrosat-intelligence.onrender.com"
        r_pl1.hyperlink.address = "https://pyrosat-intelligence.onrender.com"
        r_pl1.font.name = "Arial"
        r_pl1.font.size = Pt(7.8)
        r_pl1.font.bold = True
        r_pl1.font.underline = True
        r_pl1.font.color.rgb = RGBColor(2, 132, 199)
        p_pl1.space_after = Pt(3)

        p_pl2 = tf_proto.add_paragraph()
        r_pl2 = p_pl2.add_run()
        r_pl2.text = "🔗 Technical Defense & Benchmark Replay: GROUND_TRUTH_DEFENSE"
        r_pl2.font.name = "Arial"
        r_pl2.font.size = Pt(7.5)
        r_pl2.font.color.rgb = RGBColor(71, 85, 105)

        # ---------------------------------------------------------------------
        # 4. CENTER COLUMN: 4-Layer Orbital-to-Ground Tech Stack
        # ---------------------------------------------------------------------
        center_x = Inches(4.38)
        center_w = Inches(4.55)

        # Title for center column
        tb_ctitle = slide2.shapes.add_textbox(center_x, Inches(1.22), center_w, Inches(0.28))
        tf_ctitle = tb_ctitle.text_frame
        tf_ctitle.margin_left = tf_ctitle.margin_right = tf_ctitle.margin_top = tf_ctitle.margin_bottom = 0
        p_ct = tf_ctitle.paragraphs[0]
        p_ct.text = "CORE ARCHITECTURE: ORBITAL-TO-GROUND PIPELINE"
        p_ct.alignment = PP_ALIGN.CENTER
        p_ct.font.name = "Arial"
        p_ct.font.size = Pt(9)
        p_ct.font.bold = True
        p_ct.font.color.rgb = RGBColor(30, 58, 138)

        layers_data = [
            {
                "num": "LAYER 1",
                "title": "Spaceborne Telemetry Ingestion",
                "badge_color": RGBColor(2, 132, 199),
                "bg_color": RGBColor(248, 250, 252),
                "border_color": RGBColor(186, 230, 253),
                "bullets": [
                    "NASA VIIRS (375m) & MODIS 1km NRT thermal radiance ingestion.",
                    "Active LEO blind-window gap fill & continuous pass monitoring."
                ]
            },
            {
                "num": "LAYER 2",
                "title": "Sub-Pixel Infrared Planck Pyrometry",
                "badge_color": RGBColor(79, 70, 229),
                "bg_color": RGBColor(238, 242, 255),
                "border_color": RGBColor(199, 210, 254),
                "bullets": [
                    "Dual-Band Inversion: Solves flame temp (Tf > 1100 K) & burning area (Af).",
                    "Eliminates pixel-smearing: extracts pinpoint fires from cold background."
                ]
            },
            {
                "num": "LAYER 3",
                "title": "2-Stage Hierarchical ML + 10m Ground Truth",
                "badge_color": RGBColor(5, 150, 105),
                "bg_color": RGBColor(236, 253, 245),
                "border_color": RGBColor(167, 243, 208),
                "bullets": [
                    "TreeSHAP-verified AI segregates permitted flaring vs explosions.",
                    "Sentinel-1/2 10m ESA WorldCover validates Built-up (Class 50) vs crops."
                ]
            },
            {
                "num": "LAYER 4",
                "title": "Atmospheric Plume & NDRF Tactical Dispatch",
                "badge_color": RGBColor(217, 119, 6),
                "bg_color": RGBColor(255, 251, 235),
                "border_color": RGBColor(253, 230, 138),
                "bullets": [
                    "Briggs buoyant rise + Pasquill-Gifford Gaussian dispersion on live wind.",
                    "Instant Incident Action Plan (IAP) with automated WhatsApp/SMS dispatch."
                ]
            }
        ]

        curr_ly_y = Inches(1.52)
        layer_h = Inches(1.15)
        for i, ly in enumerate(layers_data):
            # Layer Card Background
            l_bg = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, center_x, curr_ly_y, center_w, layer_h)
            l_bg.fill.solid(); l_bg.fill.fore_color.rgb = ly["bg_color"]
            l_bg.line.color.rgb = ly["border_color"]
            l_bg.line.width = Pt(1.2)

            # Left accent tag
            l_bar = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, center_x + Inches(0.08), curr_ly_y + Inches(0.10), Inches(0.85), Inches(0.26))
            l_bar.fill.solid(); l_bar.fill.fore_color.rgb = ly["badge_color"]
            l_bar.line.fill.background()
            p_bar = l_bar.text_frame.paragraphs[0]
            p_bar.text = ly["num"]
            p_bar.alignment = PP_ALIGN.CENTER
            p_bar.font.name = "Arial"
            p_bar.font.size = Pt(7.5)
            p_bar.font.bold = True
            p_bar.font.color.rgb = WHITE

            # Layer Title
            tb_lt = slide2.shapes.add_textbox(center_x + Inches(1.02), curr_ly_y + Inches(0.08), center_w - Inches(1.12), Inches(0.30))
            tf_lt = tb_lt.text_frame; tf_lt.margin_left = tf_lt.margin_right = tf_lt.margin_top = tf_lt.margin_bottom = 0
            p_lt = tf_lt.paragraphs[0]
            p_lt.text = ly["title"]
            p_lt.font.name = "Arial"
            p_lt.font.size = Pt(9)
            p_lt.font.bold = True
            p_lt.font.color.rgb = DARK_NAVY

            # Bullets
            tb_lb = slide2.shapes.add_textbox(center_x + Inches(0.16), curr_ly_y + Inches(0.42), center_w - Inches(0.32), layer_h - Inches(0.48))
            tf_lb = tb_lb.text_frame
            tf_lb.word_wrap = True
            tf_lb.margin_left = tf_lb.margin_right = tf_lb.margin_top = tf_lb.margin_bottom = 0

            for j, b in enumerate(ly["bullets"]):
                p_b = tf_lb.paragraphs[0] if j == 0 else tf_lb.add_paragraph()
                p_b.text = f"• {b}"
                p_b.font.name = "Arial"
                p_b.font.size = Pt(7.5)
                p_b.font.color.rgb = TEXT_CHARCOAL
                if j == 0:
                    p_b.space_after = Pt(2)

            # Draw connector arrow to next layer (except last)
            if i < len(layers_data) - 1:
                arrow = slide2.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, center_x + center_w / 2 - Inches(0.10), curr_ly_y + layer_h + Inches(0.01), Inches(0.20), Inches(0.13))
                arrow.fill.solid(); arrow.fill.fore_color.rgb = RGBColor(148, 163, 184)
                arrow.line.fill.background()

            curr_ly_y += layer_h + Inches(0.15)

        # ---------------------------------------------------------------------
        # 5. RIGHT COLUMN: Risk vs Solution Matrix + Tactical Output Callout
        # ---------------------------------------------------------------------
        right_x = Inches(9.13)
        right_w = Inches(3.80)

        # Title for right column
        tb_rtitle = slide2.shapes.add_textbox(right_x, Inches(1.22), right_w, Inches(0.28))
        tf_rtitle = tb_rtitle.text_frame
        tf_rtitle.margin_left = tf_rtitle.margin_right = tf_rtitle.margin_top = tf_rtitle.margin_bottom = 0
        p_rt = tf_rtitle.paragraphs[0]
        p_rt.text = "CRITICAL RISK   vs   PYROSAT SOLUTION"
        p_rt.alignment = PP_ALIGN.CENTER
        p_rt.font.name = "Arial"
        p_rt.font.size = Pt(9)
        p_rt.font.bold = True
        p_rt.font.color.rgb = RGBColor(153, 27, 27)

        pairs_data = [
            ("Coarse 375m Satellite Pixels", "Sub-Pixel Planck Pyrometry (Tf, Af)"),
            ("Routine Flares Mistaken as Fires", "2-Stage Hierarchical ML Shield"),
            ("Blind Toxic Gas Cloud Inhalation", "3D Gaussian Plume Standoff Rings"),
            ("Delayed Multi-Agency Dispatch", "1-Click Automated NDRF Action Plan")
        ]

        curr_pair_y = Inches(1.52)
        pair_h = Inches(0.68)
        pill_w = (right_w - Inches(0.36)) / 2.0  # ~1.72 in each

        for risk, sol in pairs_data:
            # Red Risk Pill
            p_red = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x, curr_pair_y, pill_w, pair_h)
            p_red.fill.solid(); p_red.fill.fore_color.rgb = RED_BG
            p_red.line.color.rgb = RED_BORDER
            p_red.line.width = Pt(1)
            tf_r = p_red.text_frame
            tf_r.word_wrap = True
            tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = tf_r.margin_bottom = Inches(0.04)
            p_rtxt = tf_r.paragraphs[0]
            p_rtxt.text = f"🔴 {risk}"
            p_rtxt.alignment = PP_ALIGN.CENTER
            p_rtxt.font.name = "Arial"
            p_rtxt.font.size = Pt(7.5)
            p_rtxt.font.bold = True
            p_rtxt.font.color.rgb = RED_TEXT

            # Connector icon ⇄
            tb_mid = slide2.shapes.add_textbox(right_x + pill_w, curr_pair_y + Inches(0.16), Inches(0.36), Inches(0.36))
            tf_mid = tb_mid.text_frame; tf_mid.margin_left = tf_mid.margin_right = tf_mid.margin_top = tf_mid.margin_bottom = 0
            p_mtxt = tf_mid.paragraphs[0]
            p_mtxt.text = "⇄"
            p_mtxt.alignment = PP_ALIGN.CENTER
            p_mtxt.font.name = "Arial"
            p_mtxt.font.size = Pt(12)
            p_mtxt.font.bold = True
            p_mtxt.font.color.rgb = TEXT_MUTED

            # Green Solution Pill
            p_green = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, right_x + pill_w + Inches(0.36), curr_pair_y, pill_w, pair_h)
            p_green.fill.solid(); p_green.fill.fore_color.rgb = GREEN_BG
            p_green.line.color.rgb = GREEN_BORDER
            p_green.line.width = Pt(1)
            tf_g = p_green.text_frame
            tf_g.word_wrap = True
            tf_g.margin_left = tf_g.margin_right = tf_g.margin_top = tf_g.margin_bottom = Inches(0.04)
            p_gtxt = tf_g.paragraphs[0]
            p_gtxt.text = f"🟢 {sol}"
            p_gtxt.alignment = PP_ALIGN.CENTER
            p_gtxt.font.name = "Arial"
            p_gtxt.font.size = Pt(7.5)
            p_gtxt.font.bold = True
            p_gtxt.font.color.rgb = GREEN_TEXT

            curr_pair_y += pair_h + Inches(0.12)

        # Tactical Visual Callout Box at bottom right
        callout_y = curr_pair_y + Inches(0.04)
        callout_h = Inches(1.95)
        
        # Check if prototype screenshot exists to insert
        plume_img = find_asset("gaussian_plume_hazard_rings.jpg")
        if os.path.exists(plume_img):
            # Insert real graphic thumbnail
            img_w = Inches(1.70)
            img_h = Inches(1.85)
            slide2.shapes.add_picture(plume_img, right_x, callout_y, width=img_w, height=img_h)

            # Caption beside image
            tb_vcap = slide2.shapes.add_textbox(right_x + img_w + Inches(0.10), callout_y, right_w - img_w - Inches(0.10), callout_h)
            tf_vcap = tb_vcap.text_frame
            tf_vcap.word_wrap = True
            tf_vcap.margin_left = tf_vcap.margin_right = tf_vcap.margin_top = tf_vcap.margin_bottom = 0
            
            p_vc1 = tf_vcap.paragraphs[0]
            p_vc1.text = "Tactical Validation:"
            p_vc1.font.name = "Arial"
            p_vc1.font.size = Pt(8.5)
            p_vc1.font.bold = True
            p_vc1.font.color.rgb = DARK_NAVY
            p_vc1.space_after = Pt(2)

            p_vc2 = tf_vcap.add_paragraph()
            p_vc2.text = "Real-time 3D Gaussian toxic plume calculated using live Open-Meteo wind vectors with automated ERPG/IDLH evacuation rings."
            p_vc2.font.name = "Arial"
            p_vc2.font.size = Pt(7.2)
            p_vc2.font.color.rgb = TEXT_CHARCOAL

        prs.save(ppt_path)
        print(f"Successfully updated Slide 2 in {ppt_path} with clean, ATS-friendly 3-column architecture!")

if __name__ == '__main__':
    build_slide2_idea_title()
